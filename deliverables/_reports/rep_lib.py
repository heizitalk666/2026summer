"""把学院的实训报告模板填成一份报告。

模板：测控系统综合实训小组报告模板.docx（个人报告模板与它逐字节相同）
用法：content_*.py 里写内容块，build_*.py 调 build(...) 生成。

保留模板自己的东西：封面版式、页眉页脚、分节符、Heading 自动编号、摘要与参考文献样式，
只替换正文。待填的个人信息（学号、专业年级、小组编号、姓名、工作比例、答辩日期）一律黄色高亮。
"""
from __future__ import annotations

import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
TEMPLATE = os.path.join(REPO, "测控系统综合实训小组报告模板.docx")
FIG = os.path.join(HERE, "fig")


def no_indent(p):
    """模板的 Normal 带 firstLineChars 首行缩进，光设 first_line_indent 压不住，要把属性清零。"""
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    ind = p._element.get_or_add_pPr().get_or_add_ind()
    ind.set(qn("w:firstLineChars"), "0")
    ind.set(qn("w:firstLine"), "0")
    return p


def _hl(run):
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def _cjk_width(s: str) -> int:
    return sum(2 if ord(c) > 0x2E80 else 1 for c in s)


def fill_underline(par, value: str):
    """封面「小组编号 ____」这类下划线段落：把值写进第一段下划线里，宽度尽量不变。"""
    runs = [r for r in par.runs if r.underline and r.text.strip() == "" and len(r.text) >= 3]
    if not runs:
        return
    target = runs[0]
    pad = max(1, len(target.text) - _cjk_width(value))
    target.text = "  " + value + " " * max(1, pad - 2)
    _hl(target)


def fill_date(par, y: str, m: str, d: str):
    """答辩日期那一行：年、月、日三个标记各自往前找一段空白下划线填进去。"""
    runs = list(par.runs)
    marks = {"年": y, "月": m, "日": d}
    for i, r in enumerate(runs):
        if r.text.strip() in marks:
            val = marks.pop(r.text.strip())
            blanks = [x for x in runs[:i] if x.underline and x.text.strip() == "" and len(x.text) >= 2]
            if blanks:
                tgt = blanks[-1]
                pad = max(1, len(tgt.text) - _cjk_width(val))
                tgt.text = " " + val + " " * pad
                _hl(tgt)


def set_text(par, text: str, highlight: bool = False):
    """改一个段落的文字，保留它原来的样式（沿用第一个 run 的格式）。"""
    if not par.runs:
        r = par.add_run(text)
    else:
        par.runs[0].text = text
        for extra in par.runs[1:]:
            extra._element.getparent().remove(extra._element)
        r = par.runs[0]
    if highlight:
        _hl(r)
    return r


def drop_after(doc, anchor_par):
    """删掉正文区：anchor（分节符所在段落）之后的所有段落与表。"""
    body = doc.element.body
    seen = False
    for child in list(body.iterchildren()):
        if child is anchor_par._element:
            seen = True
            continue
        if seen and child.tag in (qn("w:p"), qn("w:tbl")):
            body.remove(child)


def _style(doc, name, fallback="Normal"):
    try:
        doc.styles[name]
        return name
    except KeyError:
        return fallback


def toc_field(doc, note: str):
    """真正的目录域。打开文档后按 Ctrl+A、F9 就按实际页码生成。"""
    p = doc.add_paragraph()
    r = p.add_run()
    beg = OxmlElement("w:fldChar")
    beg.set(qn("w:fldCharType"), "begin")
    r._r.append(beg)
    r2 = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-3" \h \z \u '
    r2._r.append(instr)
    r3 = p.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)
    _hl(p.add_run(note))
    r5 = p.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r5._r.append(end)
    return p


def _runs(p, text):
    """**粗体**，【】里的内容自动高亮成待填项。"""
    for seg in re.split(r"(\*\*[^*]+\*\*|【[^】]+】)", text):
        if not seg:
            continue
        if seg.startswith("**"):
            p.add_run(seg[2:-2]).bold = True
        elif seg.startswith("【"):
            _hl(p.add_run(seg))
        else:
            p.add_run(seg)


def _cell(cell, text, bold=False):
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    no_indent(p)
    for i, seg in enumerate(text.split("\n")):
        r = p.add_run(("\n" if i else "") + seg)
        r.bold = bold or seg.startswith("**")
        if seg.startswith("**"):
            r.text = r.text.replace("**", "")
        r.font.size = Pt(10.5)
        r.font.name = "宋体"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if "【" in seg:
            _hl(r)


def add_blocks(doc, blocks):
    body_style = _style(doc, "论文正文")
    cap_style = _style(doc, "Caption")
    for blk in blocks:
        kind = blk[0]
        if kind in ("h1", "h2", "h3"):
            doc.add_paragraph(blk[1], style="Heading %s" % kind[1])
        elif kind == "sty":
            doc.add_paragraph(blk[2], style=_style(doc, blk[1]))
        elif kind == "hx":
            p = no_indent(doc.add_paragraph())
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(blk[1])
            r.bold = True
            r.font.size = Pt(12)
            r.font.name = "黑体"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        elif kind == "p":
            _runs(doc.add_paragraph(style=body_style), blk[1])
        elif kind == "li":
            p = no_indent(doc.add_paragraph(style="List Paragraph"))
            p.paragraph_format.left_indent = Inches(0.35)
            _runs(p, "· " + blk[1])
        elif kind == "ref":
            doc.add_paragraph(blk[1], style=_style(doc, "参考文献"))
        elif kind == "code":
            p = no_indent(doc.add_paragraph())
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(blk[1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        elif kind == "tbl":
            _, caption, header, rows = blk
            cp = no_indent(doc.add_paragraph(caption, style=cap_style))
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            hdr = t.rows[0]._tr.get_or_add_trPr()          # 跨页时表头行重复
            rep = OxmlElement("w:tblHeader")
            hdr.append(rep)
            for c, txt in zip(t.rows[0].cells, header):
                _cell(c, txt, bold=True)
            for row in rows:
                for c, txt in zip(t.add_row().cells, row):
                    _cell(c, str(txt))
            doc.add_paragraph(style=body_style)
        elif kind == "fig":
            _, path, caption, width = blk
            p = no_indent(doc.add_paragraph())
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.keep_with_next = True       # 图和图题不拆到两页
            p.add_run().add_picture(path, width=Inches(width))
            cp = no_indent(doc.add_paragraph(caption, style=cap_style))
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            raise ValueError("未知块类型 %r" % kind)


def build(out_path, *, title_lines, cover, members, abstract_cn, keywords_cn,
          title_en, abstract_en, keywords_en, blocks):
    doc = Document(TEMPLATE)
    paras = doc.paragraphs

    # ---- 封面
    set_text(paras[5], cover["kind"])
    for i, line in enumerate(title_lines):
        p = paras[10 + i]
        r = p.add_run(line)
        r.font.size = Pt(22)
        r.font.name = "黑体"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    fill_underline(paras[15], cover["group_no"])
    fill_underline(paras[16], cover["leader"])
    fill_underline(paras[17], cover["major"])
    fill_date(paras[18], *cover["date"])

    # ---- 小组成员名单
    t = doc.tables[1]
    for i, m in enumerate(members, start=1):
        for c, txt in zip(t.rows[i].cells, [str(i)] + list(m)):
            _cell(c, txt)

    # ---- 摘要
    set_text(paras[34], "".join(title_lines))
    set_text(paras[36], abstract_cn)
    set_text(paras[38], "关键词：" + "；".join(keywords_cn))
    ep = paras[48]
    r = ep.add_run(title_en)
    r.bold = True
    r.font.size = Pt(14)
    ep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = doc.tables[2].rows[0].cells[0]
    set_text(cell.paragraphs[0], abstract_en)
    set_text(cell.paragraphs[1], "Keywords: " + "; ".join(keywords_en))
    for extra in cell.paragraphs[2:]:
        extra._element.getparent().remove(extra._element)

    # 页眉：模板里是「论文题目(五号字)」这样的占位
    header_text = "".join(title_lines)
    for sec in doc.sections:
        for hf in (sec.header, sec.first_page_header, sec.even_page_header):
            for par in hf.paragraphs:
                if "论文题目" in par.text or "题目" == par.text.strip():
                    set_text(par, header_text)

    # ---- 目录换成目录域
    for i in range(53, 69):
        el = paras[i]._element
        el.getparent().remove(el)
    anchor = doc.paragraphs[52]._element                      # 「目  录」标题
    anchor.addnext(toc_field(doc, "目录在打开文档后生成：全选（Ctrl+A）再按 F9 更新域")._element)

    # ---- 正文
    sect = [p for p in doc.paragraphs
            if p._element.find(qn("w:pPr")) is not None
            and p._element.pPr.find(qn("w:sectPr")) is not None][0]
    drop_after(doc, sect)
    add_blocks(doc, blocks)
    doc.save(out_path)
    print("已生成 %s（%.0f KB）" % (out_path, os.path.getsize(out_path) / 1024))

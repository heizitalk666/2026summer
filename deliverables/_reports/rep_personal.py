"""按学院的「个人报告模板」生成个人报告。

这份模板和小组报告模板结构不同，不能共用 rep_lib.build：
- 封面填学生姓名、学号、指导教师、专业年级、答辩日期（院系一栏模板已印好）
- 第 2 页是教师填的评分表，只填「设计题目」一行
- 第 3 页是小组成员名单
- 没有摘要；正文四章：个人工作概述 / 个人负责设计单元设计 / 团队协作与个人贡献 / 学习收获与改进建议
- 附录两样必交材料：会议记录及纪要、进度安排及调整记录

待填的个人信息一律黄色高亮。
"""
from __future__ import annotations

import os

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

from rep_lib import REPO, _cell, _cjk_width, _hl, add_blocks, drop_after, fill_date, toc_field

TEMPLATE = os.path.join(REPO, "测控系统综合实训个人报告模板.docx")


def _fill_segment(par, label: str, value: str, stop: str | None = None):
    """一行里有几栏（「学生姓名 ___ 学号 ___」）：把 value 写进 label 之后、stop 之前最长的那段下划线。"""
    runs = list(par.runs)
    start = next(i for i, r in enumerate(runs) if r.text.strip() == label) + 1
    end = next((i for i, r in enumerate(runs) if i >= start and stop and r.text.strip() == stop), len(runs))
    blanks = [r for r in runs[start:end] if r.underline and r.text.strip() == "" and "\t" not in r.text]
    if not blanks:
        return
    tgt = max(blanks, key=lambda r: len(r.text))
    pad = max(2, len(tgt.text) - _cjk_width(value))
    tgt.text = " " * (pad // 2) + value + " " * (pad - pad // 2)
    if value.startswith("【"):
        _hl(tgt)


def _replace_runs(par, old_texts: list[str], new: str):
    """把相邻的几个 run（如 "20" "xx" "级自动化"）合成一个新值，保留第一个 run 的格式。"""
    runs = [r for r in par.runs if r.text in old_texts]
    if not runs:
        return
    runs[0].text = new
    _hl(runs[0])
    for r in runs[1:]:
        r.text = ""


def build(out_path, *, title_lines, name, student_id, advisor, major, date, members, blocks):
    doc = Document(TEMPLATE)
    paras = doc.paragraphs

    # 正文样式：模板正文段是 Normal 加两字首行缩进，rep_lib.add_blocks 找的是「论文正文」
    st = doc.styles.add_style("论文正文", WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles["Normal"]
    st.paragraph_format.first_line_indent = Emu(304800)

    # ---- 封面
    for i, line in enumerate(title_lines):
        p = paras[8 + i]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        r = p.add_run(line)
        r.font.size = Pt(22)
        r.font.name = "黑体"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    _fill_segment(paras[10], "学生姓名", name, stop="学号")
    _fill_segment(paras[10], "学号", student_id)
    _fill_segment(paras[11], "指导教师", advisor)
    _replace_runs(paras[13], ["20", "xx", "级自动化"], major)
    fill_date(paras[14], *date)

    # ---- 评分表：只填设计题目，其余由教师填
    _cell(doc.tables[1].rows[0].cells[1], "".join(title_lines))

    # ---- 小组成员名单
    t = doc.tables[2]
    for i, m in enumerate(members, start=1):
        for c, txt in zip(t.rows[i].cells, [str(i)] + list(m)):
            _cell(c, txt)

    # ---- 目录：模板里是写死的条目，换成目录域
    toc_title = next(p for p in paras if p.text.replace(" ", "") == "目录")
    k = paras.index(toc_title)
    for p in paras[k + 1:]:
        if p.style.name.startswith("toc"):
            p._element.getparent().remove(p._element)
    toc_title._element.addnext(toc_field(doc, "目录在打开文档后生成：全选（Ctrl+A）再按 F9 更新域")._element)

    # ---- 正文
    sect = [p for p in doc.paragraphs
            if p._element.find(qn("w:pPr")) is not None
            and p._element.pPr.find(qn("w:sectPr")) is not None][0]
    drop_after(doc, sect)
    n_tpl_tables = len(doc.tables)
    add_blocks(doc, blocks)

    # 模板 Normal 是 1.25 倍行距，正文合适，表格和代码段用单倍，不然一张表要多占半页
    for t in doc.tables[n_tpl_tables:]:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
    for p in doc.paragraphs:
        if p.runs and p.runs[0].font.name == "Consolas":
            p.paragraph_format.line_spacing = 1.0
    doc.save(out_path)
    print("已生成 %s（%.0f KB）" % (out_path, os.path.getsize(out_path) / 1024))
